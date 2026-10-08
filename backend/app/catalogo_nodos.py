from pathlib import Path
import json
from openpyxl import load_workbook

from .config import BASE_DIR


ARCHIVO_MAESTRO_NODOS = (
    BASE_DIR.parent
    / "Servicio Sentido Terminal Nodo Alfa Omega.xlsx"
)

CARPETA_CONFIGURACION = (
    BASE_DIR
    / "configuracion"
)

ARCHIVO_CATALOGO_NODOS = (
    CARPETA_CONFIGURACION
    / "catalogo_nodos.json"
)


HOJAS_UNIDADES = {
    "SS T Alfa": "ALFA",
    "SS T Omega": "OMEGA",
}


def normalizar_texto(valor):
    if valor is None:
        return ""

    return " ".join(
        str(valor).strip().split()
    )


def cargar_catalogo_desde_excel():

    if not ARCHIVO_MAESTRO_NODOS.exists():
        raise FileNotFoundError(
            f"No existe catálogo maestro: "
            f"{ARCHIVO_MAESTRO_NODOS}"
        )

    wb = load_workbook(
        ARCHIVO_MAESTRO_NODOS,
        read_only=True,
        data_only=True,
    )

    registros = []

    for hoja, unidad in HOJAS_UNIDADES.items():

        if hoja not in wb.sheetnames:
            continue

        ws = wb[hoja]

        for fila in ws.iter_rows(
            values_only=True
        ):

            codigo = (
                fila[1]
                if len(fila) > 1
                else None
            )

            terminal = (
                fila[2]
                if len(fila) > 2
                else None
            )

            nodo = (
                fila[3]
                if len(fila) > 3
                else None
            )

            codigo = normalizar_texto(
                codigo
            )

            if not codigo:
                continue

            if codigo.upper() in (
                "SERVICIO",
                "SERVICIOS SENTIDOS POR TERMINAL",
            ):
                continue

            registros.append({
                "unidad": unidad,
                "codigo": codigo,
                "terminal": normalizar_texto(
                    terminal
                ),
                "nodo": normalizar_texto(
                    nodo
                ),
            })

    wb.close()

    registros.sort(
        key=lambda r: (
            r["unidad"],
            r["codigo"],
        )
    )

    return registros


def guardar_catalogo(registros):

    CARPETA_CONFIGURACION.mkdir(
        parents=True,
        exist_ok=True,
    )

    temporal = (
        ARCHIVO_CATALOGO_NODOS
        .with_suffix(".tmp")
    )

    contenido = {
        "version": 1,
        "registros": registros,
    }

    temporal.write_text(
        json.dumps(
            contenido,
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )

    temporal.replace(
        ARCHIVO_CATALOGO_NODOS
    )


def asegurar_catalogo():

    if not ARCHIVO_CATALOGO_NODOS.exists():

        registros = (
            cargar_catalogo_desde_excel()
        )

        guardar_catalogo(
            registros
        )

    return ARCHIVO_CATALOGO_NODOS


def obtener_catalogo():

    asegurar_catalogo()

    datos = json.loads(
        ARCHIVO_CATALOGO_NODOS.read_text(
            encoding="utf-8"
        )
    )

    return datos.get(
        "registros",
        []
    )


def actualizar_nodo(
    unidad,
    codigo,
    nodo,
):

    unidad = normalizar_texto(
        unidad
    ).upper()

    codigo = normalizar_texto(
        codigo
    )

    nodo = normalizar_texto(
        nodo
    )

    registros = obtener_catalogo()

    encontrado = False

    for registro in registros:

        if (
            registro["unidad"].upper()
            == unidad
            and
            registro["codigo"].upper()
            == codigo.upper()
        ):
            registro["nodo"] = nodo
            encontrado = True
            break

    if not encontrado:
        raise ValueError(
            "No existe la combinación "
            f"{unidad} / {codigo}"
        )

    guardar_catalogo(
        registros
    )

    return registro
